"""Build FUNDING_PROPOSAL.docx from FUNDING_PROPOSAL.md.

Hand-rolled markdown -> docx so the output has clean styling for a professor:
- H1/H2/H3 headings
- Real Word tables for the budget tables
- Proper bullet lists
- Bold/italic inline formatting
- Justified body paragraphs
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "FUNDING_PROPOSAL.md"
DST = ROOT / "FUNDING_PROPOSAL.docx"


INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
INLINE_CODE = re.compile(r"`([^`]+)`")
INLINE_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_runs(paragraph, text: str) -> None:
    """Parse a single line of inline markdown into runs on a paragraph."""
    # Strip links to just their visible text
    text = INLINE_LINK.sub(lambda m: m.group(1), text)

    # Tokenize by bold/italic/code in one pass
    pattern = re.compile(
        r"(\*\*[^*]+\*\*)"        # bold
        r"|(\*[^*]+\*)"           # italic
        r"|(`[^`]+`)"             # code
    )
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def build():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()

    # Default style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()

        # Skip blank lines
        if not line.strip():
            i += 1
            continue

        # Horizontal rule -> visual separator
        if line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("—" * 30)
            run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Headings
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue

        # Tables (GFM)
        if line.lstrip().startswith("|") and i + 1 < n and re.match(r"\s*\|[\s\-:|]+\|\s*$", lines[i + 1]):
            header = parse_table_row(line)
            i += 2  # skip separator
            body_rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                body_rows.append(parse_table_row(lines[i]))
                i += 1

            table = doc.add_table(rows=1 + len(body_rows), cols=len(header))
            table.style = "Light Grid Accent 1"

            # Header
            for col, h_text in enumerate(header):
                cell = table.rows[0].cells[col]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(h_text)
                run.bold = True

            # Body
            for r, row in enumerate(body_rows, start=1):
                for col, cell_text in enumerate(row):
                    if col >= len(header):
                        continue
                    cell = table.rows[r].cells[col]
                    cell.text = ""
                    add_runs(cell.paragraphs[0], cell_text)

            doc.add_paragraph()  # spacing after table
            continue

        # Bulleted list
        if line.lstrip().startswith(("- ", "* ")):
            content = line.lstrip()[2:]
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, content)
            i += 1
            continue

        # Numbered list
        m_num = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m_num:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m_num.group(1))
            i += 1
            continue

        # Blockquote
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run_marker = p.add_run("“")
            run_marker.italic = True
            add_runs(p, line[2:])
            i += 1
            continue

        # Regular paragraph — collect continuation lines until blank
        buf = [line]
        i += 1
        while i < n and lines[i].strip() and not re.match(
            r"^(#|\||-\s|\*\s|>\s|\d+\.\s|---)", lines[i].lstrip()
        ):
            buf.append(lines[i].rstrip())
            i += 1

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, " ".join(buf))

    doc.save(DST)
    print(f"Wrote {DST}")


if __name__ == "__main__":
    build()
